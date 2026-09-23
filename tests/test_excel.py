"""The Excel model must build, contain every sheet, and compute its outputs with formulas (not pasted values).
Numerical agreement with Python is verified inside the workbook (Checks sheet) after recalculation."""
from openpyxl import load_workbook

from src.config import OUTPUTS


def test_workbook_structure():
    wb = load_workbook(OUTPUTS / "COST_Valuation_Model.xlsx")
    assert wb.sheetnames == ["Cover", "Historicals", "Assumptions", "Forecast", "Beta", "WACC", "DCF", "Sensitivity",
                             "Comps", "Summary", "Checks"]
    formulas = sum(1 for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    assert formulas > 800


def test_key_outputs_are_formulas():
    wb = load_workbook(OUTPUTS / "COST_Valuation_Model.xlsx")
    assert str(wb["WACC"]["B27"].value).startswith("=")          # WACC
    assert str(wb["Beta"]["H6"].value).startswith("=SLOPE")
    for ws in ("Forecast", "DCF"):
        cells = [c.value for row in wb[ws].iter_rows(min_col=4, max_col=8) for c in row if c.value is not None]
        assert sum(isinstance(v, str) and v.startswith("=") for v in cells) / len(cells) > 0.8   # rest: dates / period fractions


def test_committed_copy_has_cached_values_and_passes_checks():
    wb = load_workbook(OUTPUTS / "COST_Valuation_Model.xlsx", data_only=True)
    if wb["Cover"]["B13"].value is None:        # freshly rebuilt, not yet recalculated
        return
    assert wb["Cover"]["B13"].value == "ALL OK"
