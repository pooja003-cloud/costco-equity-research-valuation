"""Builds outputs/COST_Valuation_Model.xlsx, a formula-driven copy of the Python model.

Sheets: Cover, Historicals, Assumptions, Forecast, Beta, WACC, DCF, Sensitivity, Comps, Summary, Checks.
Blue = input, black = formula, green = link to another sheet, yellow = key assumption.
Open in Excel or LibreOffice afterwards so the formulas recalculate.
"""
from __future__ import annotations

import contextlib
import io
import json
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

from openpyxl.utils import get_column_letter as L

from src.config import CONFIG_DIR, OUTPUTS, PROCESSED, TABLES
from src.market import prices

FONT = "Arial"
BLUE, GREEN, BLACK, GREY = "0000FF", "008000", "000000", "666666"
F_IN = Font(name=FONT, size=10, color=BLUE)
F_LINK = Font(name=FONT, size=10, color=GREEN)
F_CALC = Font(name=FONT, size=10, color=BLACK)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color=GREY)
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
FILL_SUB = PatternFill("solid", fgColor="D9E1F2")

NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT2 = '0.00%;(0.00%);"-"'
MULT = '0.0"x";(0.0"x");"-"'
DEC = '0.000'
DATE = 'dd-mmm-yyyy'

HIST_YEARS = [2019, 2020, 2021, 2022, 2023, 2024]
FC_YEARS = [2025, 2026, 2027, 2028, 2029]


def fin_long():
    return pd.read_csv(PROCESSED / "financials_long.csv")


class S:
    """Thin wrapper around a worksheet that remembers where each labelled row lives."""

    def __init__(self, ws):
        self.ws, self.rows = ws, {}

    def cell(self, ref, value=None, font=F_CALC, fmt=None, fill=None, bold=False, comment=None, align=None):
        c = self.ws[ref]
        if value is not None:
            c.value = value
        c.font = Font(name=FONT, size=font.size, color=font.color, bold=bold or font.bold, italic=font.italic)
        if fmt:
            c.number_format = fmt
        if fill:
            c.fill = fill
        if comment:
            c.comment = Comment(comment, "Pooja Master")
        if align:
            c.alignment = Alignment(horizontal=align)
        return c

    def header(self, row, labels, start_col=1):
        for i, lab in enumerate(labels):
            c = self.cell(f"{L(start_col + i)}{row}", lab, font=F_HEAD, fill=FILL_HEAD)
            c.alignment = Alignment(horizontal="center" if i else "left")

    def title(self, text, sub=None):
        self.cell("A1", text, font=F_TITLE)
        self.cell("A2", "Independent academic research and valuation case study; not investment advice.", font=F_NOTE)
        if sub:
            self.cell("A3", sub, font=F_NOTE)

    def widths(self, **w):
        for k, v in w.items():
            self.ws.column_dimensions[k].width = v


def ref(sheet, cell, absolute=True):
    if absolute:
        col = "".join(ch for ch in cell if ch.isalpha())
        row = "".join(ch for ch in cell if ch.isdigit())
        cell = f"${col}${row}"
    return f"'{sheet}'!{cell}" if " " in sheet else f"{sheet}!{cell}"


def fill_placeholders(text, rows):
    """Replace {key} placeholders with row numbers (longest keys first so {ebit} does not clobber {ebitda})."""
    if not isinstance(text, str):
        return text
    for k in sorted(rows, key=len, reverse=True):
        text = text.replace("{" + k + "}", str(rows[k]))
    return text


def build_historicals(wb):
    s = S(wb.create_sheet("Historicals"))
    s.title("Historical financial statements, FY2019-FY2024 ($ millions)",
            "Source: Costco 10-K inline XBRL via src/extract_financials.py. Each value has a comment with its XBRL concept, "
            "filing accession and a link to the tagged fact.")
    fl = fin_long()
    cols = {y: L(2 + i) for i, y in enumerate(HIST_YEARS)}
    s.header(5, ["$ millions, fiscal years ending ~31 Aug"] + [f"FY{y}" for y in HIST_YEARS])
    r = 6
    for title, st in [("Income statement", "income_statement"), ("Balance sheet", "balance_sheet"),
                      ("Cash flow", "cash_flow"), ("Lease note", "lease_note")]:
        s.cell(f"A{r}", title, font=F_BOLD, fill=FILL_SUB)
        for y in HIST_YEARS:
            s.cell(f"{cols[y]}{r}", None, fill=FILL_SUB)
        r += 1
        d = fl[fl.statement == st]
        for key in dict.fromkeys(d.line_item):
            rows = d[d.line_item == key]
            s.cell(f"A{r}", rows.label.iloc[0])
            for _, x in rows.iterrows():
                y = int(x.fiscal_year)
                fmt = USD2 if key.startswith("eps") else (NUM1 if "shares" in key else NUM)
                if isinstance(x.concept, str):
                    dims = f" [{x.dims}]" if isinstance(x.dims, str) and x.dims else ""
                    note = f"XBRL {x.concept}{dims}\nAccession {x.accession}\n{x.source_url}"
                else:
                    note = str(x.note)
                s.cell(f"{cols[y]}{r}", float(x.value), font=F_IN, fmt=fmt, comment=note)
            s.rows[key] = r
            r += 1
        r += 1
    om = pd.read_csv(PROCESSED / "operating_metrics_long.csv")
    s.cell(f"A{r}", "Operating metrics", font=F_BOLD, fill=FILL_SUB)
    r += 1
    for key, lab, fmt in [("warehouses_end_of_year", "Warehouses (end of year)", NUM), ("paid_members", "Paid members (m)", NUM1),
                          ("executive_members", "Executive members (m)", NUM1),
                          ("comparable_sales_growth", "Comparable sales growth (%)", NUM1),
                          ("renewal_rate_us_canada", "Renewal rate US & Canada (%)", NUM1)]:
        s.cell(f"A{r}", lab)
        for _, x in om[om.metric == key].iterrows():
            y = int(x.fiscal_year)
            if y in cols:
                src = f'Quote: "{x.quote}"' if isinstance(x.quote, str) and x.quote else str(x.source_url)
                s.cell(f"{cols[y]}{r}", float(x.value), font=F_IN, fmt=fmt,
                       comment=f"{x.section}\nAccession {x.source_accession}\n{src}")
        s.rows[key] = r
        r += 1
    r += 1
    s.cell(f"A{r}", "Ratios (formulas)", font=F_BOLD, fill=FILL_SUB)
    r += 1
    R = s.rows
    ratios = [
        ("Revenue growth", lambda c, p: f"={c}{R['total_revenue']}/{p}{R['total_revenue']}-1", PCT),
        ("Gross margin on net sales", lambda c, p: f"=1-{c}{R['merchandise_costs']}/{c}{R['net_sales']}", PCT2),
        ("SG&A % net sales", lambda c, p: f"={c}{R['sga']}/{c}{R['net_sales']}", PCT2),
        ("Operating margin", lambda c, p: f"={c}{R['operating_income']}/{c}{R['total_revenue']}", PCT2),
        ("Membership fees / operating income", lambda c, p: f"={c}{R['membership_fees']}/{c}{R['operating_income']}", PCT),
        ("Effective tax rate", lambda c, p: f"={c}{R['income_tax']}/{c}{R['pretax_income']}", PCT),
        ("Capex % revenue", lambda c, p: f"={c}{R['capex']}/{c}{R['total_revenue']}", PCT2),
        ("Free cash flow (CFO - capex)", lambda c, p: f"={c}{R['cfo']}-{c}{R['capex']}", NUM),
    ]
    for lab, f, fmt in ratios:
        s.cell(f"A{r}", lab)
        for i, y in enumerate(HIST_YEARS[1:], start=1):
            if lab == "Revenue growth" and y == 2020:
                continue  # FY2019 income statement is outside the dataset (only its balance sheet is used)
            s.cell(f"{cols[y]}{r}", f(cols[y], cols[HIST_YEARS[i - 1]]), fmt=fmt)
        r += 1
    s.widths(A=58, **{cols[y]: 12 for y in HIST_YEARS})
    s.ws.freeze_panes = "B6"
    return s, cols


def build_assumptions(wb):
    s = S(wb.create_sheet("Assumptions"))
    s.title("Forecast and valuation assumptions (FY2025-FY2029)",
            "Values and rationale mirror config/assumptions.json, config/valuation.json and config/scenarios.json. "
            "Change blue cells; the model updates.")
    a = json.load(open(CONFIG_DIR / "assumptions.json"))
    v = json.load(open(CONFIG_DIR / "valuation.json"))
    sc = json.load(open(CONFIG_DIR / "scenarios.json"))
    s.cell("A5", "Scenario selector (1 = Bear, 2 = Base, 3 = Bull)", bold=True)
    s.cell("C5", 2, font=F_IN, fill=FILL_KEY, comment="Change to 1 or 3 to run the bear or bull case. Base = 2.")
    s.cell("D5", '=CHOOSE($C$5,"Bear","Base","Bull")', bold=True)
    yc = {y: L(3 + i) for i, y in enumerate(FC_YEARS)}          # C..G
    s.header(7, ["Driver", "Unit"] + [f"FY{y}E" for y in FC_YEARS] + ["Rationale / source"])
    r = 8

    def scen_block(key, label, unit, fmt, base_vals, bear_vals, bull_vals, rationale):
        nonlocal r
        s.cell(f"A{r}", label, bold=True)
        s.cell(f"B{r}", unit)
        active, blocks = r, {}
        for off, name, vals in [(1, "Bear", bear_vals), (2, "Base", base_vals), (3, "Bull", bull_vals)]:
            rr = r + off
            s.cell(f"A{rr}", f"   {name}", font=F_NOTE)
            for y in FC_YEARS:
                s.cell(f"{yc[y]}{rr}", vals[y], font=F_IN, fmt=fmt)
            blocks[name] = rr
        for y in FC_YEARS:
            c = yc[y]
            s.cell(f"{c}{active}", f"=CHOOSE($C$5,{c}{blocks['Bear']},{c}{blocks['Base']},{c}{blocks['Bull']})", fmt=fmt, fill=FILL_KEY)
        s.cell(f"H{active}", rationale, font=F_NOTE)
        s.rows[key] = active
        r += 5

    def single(key, label, unit, fmt, vals, rationale):
        nonlocal r
        s.cell(f"A{r}", label, bold=True)
        s.cell(f"B{r}", unit)
        for y in FC_YEARS:
            s.cell(f"{yc[y]}{r}", vals[y], font=F_IN, fmt=fmt)
        s.cell(f"H{r}", rationale, font=F_NOTE)
        s.rows[key] = r
        r += 1

    yv = lambda d: {y: d[str(y)] for y in FC_YEARS}
    const = lambda x: {y: x for y in FC_YEARS}
    scen_block("net_new", "Net new warehouses", "count", NUM, yv(a["warehouses"]["net_new"]),
               yv(sc["bear"]["net_new_warehouses"]), yv(sc["bull"]["net_new_warehouses"]), a["warehouses"]["rationale"])
    scen_block("comps", "Comparable sales growth (reported)", "%", PCT, yv(a["comparable_sales_growth"]["values"]),
               yv(sc["bear"]["comparable_sales_growth"]), yv(sc["bull"]["comparable_sales_growth"]), a["comparable_sales_growth"]["rationale"])
    scen_block("member_growth", "Paid member growth", "%", PCT, yv(a["paid_member_growth"]["values"]),
               yv(sc["bear"]["paid_member_growth"]), yv(sc["bull"]["paid_member_growth"]), a["paid_member_growth"]["rationale"])
    scen_block("gm", "Gross margin on net sales", "%", PCT2, yv(a["gross_margin_on_net_sales"]["values"]),
               const(sc["bear"]["gross_margin_on_net_sales"]), const(sc["bull"]["gross_margin_on_net_sales"]),
               a["gross_margin_on_net_sales"]["rationale"])
    scen_block("sga", "SG&A % of net sales", "%", PCT2, yv(a["sga_pct_net_sales"]["values"]),
               const(sc["bear"]["sga_pct_net_sales"]), const(sc["bull"]["sga_pct_net_sales"]), a["sga_pct_net_sales"]["rationale"])
    r += 1
    s.cell(f"A{r}", "Single-value drivers", font=F_BOLD, fill=FILL_SUB)
    r += 1
    fi = a["membership_fee_increase"]
    single("productivity", "New-warehouse productivity (x unit growth)", "x", DEC, const(a["new_warehouse_productivity"]["value"]),
           a["new_warehouse_productivity"]["rationale"])
    single("fee_uplift", "Fee increase on affected fees", "%", PCT, const(fi["full_uplift_on_affected_fees"]),
           "$60 to $65 and $120 to $130 (+8.3%), FY2024 10-K")
    single("fee_share", "Share of fee income affected (U.S. & Canada)", "%", PCT, const(fi["share_of_fee_income_affected"]), fi["rationale"])
    single("fee_recog", "Share of fee increase recognised in year", "%", PCT,
           {y: fi["recognition_by_year"].get(str(y), 0.0) for y in FC_YEARS},
           "Fees are recognised ratably over the 12-month membership; renewals are spread through the year")
    single("fee_other", "Other fee-per-member growth", "%", PCT, const(a["fee_per_member_other_growth"]["value"]),
           a["fee_per_member_other_growth"]["rationale"])
    single("tax", "Tax rate", "%", PCT, const(a["tax_rate"]["value"]), a["tax_rate"]["rationale"])
    single("da_pct", "D&A % of beginning net PP&E", "%", PCT, const(a["d_and_a_pct_beginning_ppe"]["value"]),
           a["d_and_a_pct_beginning_ppe"]["rationale"])
    single("capex_pct", "Capex % of revenue (FY2026+)", "%", PCT2, const(a["capex"]["pct_revenue_after"]), a["capex"]["rationale"])
    s.cell(f"A{r}", "Capex FY2025 (management plan)", bold=True)
    s.cell(f"B{r}", "$m")
    s.cell(f"C{r}", a["capex"]["fy2025_dollars"], font=F_IN, fmt=NUM,
           comment="Q1 FY2025 10-Q: 'it is our current intention to spend a total of approximately $5,000 during fiscal 2025'")
    s.rows["capex_fy25"] = r
    r += 1
    wc = a["working_capital_pct_revenue"]
    for k, lab in [("receivables", "Receivables % revenue"), ("inventories", "Inventories % revenue"),
                   ("other_current_assets", "Other current assets % revenue"), ("accounts_payable", "Accounts payable % revenue"),
                   ("accrued_salaries", "Accrued salaries % revenue"), ("accrued_member_rewards", "Accrued member rewards % revenue"),
                   ("other_current_liabilities", "Other current liabilities % revenue")]:
        single(f"wc_{k}", lab, "%", PCT2, const(wc[k]), "Held at FY2024 ratio")
    single("wc_deferred_fees", "Deferred membership fees % fee income", "%", PCT,
           const(wc["deferred_membership_fees_pct_fee_income"]), "Held at FY2024 ratio")
    ii = a["interest"]
    single("int_exp", "Interest expense", "$m", NUM, const(ii["expense_fixed"]), ii["rationale"])
    single("int_yield", "Interest income yield on beginning cash + STI", "%", PCT2, yv(ii["income_yield_on_beginning_cash"]), ii["rationale"])
    cs = a["capital_structure"]
    single("payout", "Dividend payout of net income", "%", PCT, const(cs["dividend_payout_of_net_income"]), cs["rationale"])
    single("buybacks", "Share repurchases", "$m", NUM, const(cs["buybacks_per_year"]), cs["rationale"])
    single("sbc_pct", "Stock-based compensation % revenue", "%", PCT2, const(cs["sbc_pct_revenue"]), "FY2024 level (0.32%)")
    single("shares", "Diluted shares", "m", NUM1, const(cs["diluted_shares_m"]), "Q1 FY2025 10-Q diluted weighted shares")
    r += 1
    s.cell(f"A{r}", "Terminal value (C = active scenario; D/E/F = Bear/Base/Bull)", font=F_BOLD, fill=FILL_SUB)
    r += 1
    for key, lab, base, bear, bull, rat in [
            ("tg", "Terminal growth (g)", v["terminal"]["growth"], sc["bear"]["terminal_growth"], sc["bull"]["terminal_growth"],
             v["terminal"]["rationale_growth"]),
            ("ronic", "Terminal return on new invested capital (RONIC)", v["terminal"]["ronic"], sc["bear"]["terminal_ronic"],
             sc["bull"]["terminal_ronic"], v["terminal"]["rationale_ronic"])]:
        s.cell(f"A{r}", lab, bold=True)
        s.cell(f"B{r}", "%")
        s.cell(f"D{r}", bear, font=F_IN, fmt=PCT)
        s.cell(f"E{r}", base, font=F_IN, fmt=PCT)
        s.cell(f"F{r}", bull, font=F_IN, fmt=PCT)
        s.cell(f"C{r}", f"=CHOOSE($C$5,D{r},E{r},F{r})", fmt=PCT, fill=FILL_KEY)
        s.cell(f"H{r}", rat, font=F_NOTE)
        s.rows[key] = r
        r += 1
    s.widths(A=48, B=7, C=11, D=11, E=11, F=11, G=11, H=120)
    s.ws.freeze_panes = "C8"
    return s, yc


def build_forecast(wb, H, hc, A, ac):
    s = S(wb.create_sheet("Forecast"))
    s.title("Integrated three-statement forecast ($ millions)",
            "FY2024A links to Historicals; FY2025E-FY2029E are formulas driven by Assumptions. Cash is the only plug; see the balance check.")
    cols = {2024: "C", **{y: L(4 + i) for i, y in enumerate(FC_YEARS)}}   # C = FY2024A, D..H forecast
    s.header(5, ["$ millions", "Note", "FY2024A"] + [f"FY{y}E" for y in FC_YEARS])
    Hr, Ar = H.rows, A.rows
    h = lambda key: ref("Historicals", f"{hc[2024]}{Hr[key]}")
    h23 = lambda key: ref("Historicals", f"{hc[2023]}{Hr[key]}")
    a = lambda key, y: ref("Assumptions", f"{ac[y]}{Ar[key]}")
    R = s.rows
    state = {"r": 6}

    def is_pure_link(f):
        return isinstance(f, str) and f.startswith("=") and f.count("!") == 1 and not any(op in f[1:] for op in "+-*/(")

    def line(key, label, base, fc, fmt=NUM, bold=False):
        r = state["r"]
        s.cell(f"A{r}", label, bold=bold)
        if base is not None:
            s.cell(f"C{r}", base, font=F_LINK if is_pure_link(base) else F_CALC, fmt=fmt, bold=bold)
        for y in FC_YEARS:
            f = fc(y, cols[y], cols[y - 1])
            s.cell(f"{cols[y]}{r}", f, font=F_LINK if is_pure_link(f) else F_CALC, fmt=fmt, bold=bold)
        R[key] = r
        state["r"] += 1

    def section(t):
        r = state["r"]
        s.cell(f"A{r}", t, font=F_BOLD, fill=FILL_SUB)
        for c in "BCDEFGH":
            s.cell(f"{c}{r}", None, fill=FILL_SUB)
        state["r"] += 1

    section("Operating drivers")
    line("warehouses", "Warehouses (end of year)", "=" + h("warehouses_end_of_year"), lambda y, c, p: f"={p}{{warehouses}}+{a('net_new', y)}")
    line("unit_growth", "Net unit growth", None, lambda y, c, p: f"={a('net_new', y)}/{p}{{warehouses}}", PCT)
    line("comps", "Comparable sales growth", None, lambda y, c, p: f"={a('comps', y)}", PCT)
    line("nuc", "Growth from new warehouses", None, lambda y, c, p: f"={a('productivity', y)}*{c}{{unit_growth}}", PCT)
    line("ns_growth", "Net sales growth", None, lambda y, c, p: f"=(1+{c}{{comps}})*(1+{c}{{nuc}})-1", PCT)
    line("members", "Paid members, end of year (m)", "=" + h("paid_members"), lambda y, c, p: f"={p}{{members}}*(1+{a('member_growth', y)})", NUM1)
    line("fee_effect", "Fee increase recognised", None, lambda y, c, p: f"={a('fee_uplift', y)}*{a('fee_share', y)}*{a('fee_recog', y)}", PCT)
    line("fpm", "Fee income per average paid member ($)", f"={h('membership_fees')}/AVERAGE({h('paid_members')},{h23('paid_members')})",
         lambda y, c, p: f"={p}{{fpm}}*(1+{c}{{fee_effect}})*(1+{a('fee_other', y)})", USD2)
    section("Income statement")
    line("net_sales", "Net sales", "=" + h("net_sales"), lambda y, c, p: f"={p}{{net_sales}}*(1+{c}{{ns_growth}})")
    line("fees", "Membership fees", "=" + h("membership_fees"), lambda y, c, p: f"={c}{{fpm}}*AVERAGE({c}{{members}},{p}{{members}})")
    line("revenue", "Total revenue", "=C{net_sales}+C{fees}", lambda y, c, p: f"={c}{{net_sales}}+{c}{{fees}}", bold=True)
    line("cogs", "Merchandise costs", "=" + h("merchandise_costs"), lambda y, c, p: f"={c}{{net_sales}}*(1-{a('gm', y)})")
    line("sga", "SG&A", "=" + h("sga"), lambda y, c, p: f"={c}{{net_sales}}*{a('sga', y)}")
    line("ebit", "Operating income (EBIT)", "=C{revenue}-C{cogs}-C{sga}", lambda y, c, p: f"={c}{{revenue}}-{c}{{cogs}}-{c}{{sga}}", bold=True)
    line("int_exp", "Interest expense", "=" + h("interest_expense"), lambda y, c, p: f"={a('int_exp', y)}")
    line("int_inc", "Interest income and other", "=" + h("interest_income_other"), lambda y, c, p: f"={a('int_yield', y)}*({p}{{cash}}+{p}{{sti}})")
    line("pretax", "Income before taxes", "=C{ebit}-C{int_exp}+C{int_inc}", lambda y, c, p: f"={c}{{ebit}}-{c}{{int_exp}}+{c}{{int_inc}}")
    line("tax", "Income taxes", "=" + h("income_tax"), lambda y, c, p: f"={c}{{pretax}}*{a('tax', y)}")
    line("ni", "Net income", "=" + h("net_income"), lambda y, c, p: f"={c}{{pretax}}-{c}{{tax}}", bold=True)
    line("shares", "Diluted shares (m)", "=" + h("shares_diluted"), lambda y, c, p: f"={a('shares', y)}", NUM1)
    line("eps", "Diluted EPS ($)", "=" + h("eps_diluted"), lambda y, c, p: f"={c}{{ni}}/{c}{{shares}}", USD2)
    line("da", "Depreciation & amortisation", "=" + h("d_and_a"), lambda y, c, p: f"={a('da_pct', y)}*{p}{{ppe}}")
    line("ebitda", "EBITDA", "=C{ebit}+C{da}", lambda y, c, p: f"={c}{{ebit}}+{c}{{da}}")
    line("opm", "Operating margin", "=C{ebit}/C{revenue}", lambda y, c, p: f"={c}{{ebit}}/{c}{{revenue}}", PCT2)
    section("Balance sheet")
    line("cash", "Cash and cash equivalents", "=" + h("cash"), lambda y, c, p: f"={p}{{cash}}+{c}{{net_change}}")
    line("sti", "Short-term investments", "=" + h("short_term_investments"), lambda y, c, p: f"={p}{{sti}}")
    for k, lab in [("receivables", "Receivables"), ("inventories", "Merchandise inventories"), ("other_current_assets", "Other current assets")]:
        line(k, lab, "=" + h(k), lambda y, c, p, k=k: f"={a('wc_' + k, y)}*{c}{{revenue}}")
    line("ppe", "Property and equipment, net", "=" + h("ppe_net"), lambda y, c, p: f"={p}{{ppe}}+{c}{{capex}}-{c}{{da}}")
    line("rou", "Operating lease ROU assets", "=" + h("operating_lease_rou"), lambda y, c, p: f"={p}{{rou}}")
    line("olta", "Other long-term assets", "=" + h("other_lt_assets"), lambda y, c, p: f"={p}{{olta}}")
    line("ta", "Total assets", "=SUM(C{cash}:C{olta})", lambda y, c, p: f"=SUM({c}{{cash}}:{c}{{olta}})", bold=True)
    for k, lab in [("accounts_payable", "Accounts payable"), ("accrued_salaries", "Accrued salaries and benefits"),
                   ("accrued_member_rewards", "Accrued member rewards"), ("other_current_liabilities", "Other current liabilities")]:
        line(k, lab, "=" + h(k), lambda y, c, p, k=k: f"={a('wc_' + k, y)}*{c}{{revenue}}")
    line("dmf", "Deferred membership fees", "=" + h("deferred_membership_fees"), lambda y, c, p: f"={a('wc_deferred_fees', y)}*{c}{{fees}}")
    line("cdebt", "Current portion of long-term debt", "=" + h("current_debt"), lambda y, c, p: f"={p}{{cdebt}}")
    line("ltdebt", "Long-term debt", "=" + h("long_term_debt"), lambda y, c, p: f"={p}{{ltdebt}}")
    line("oll", "Long-term operating lease liabilities", "=" + h("lt_operating_lease_liab"), lambda y, c, p: f"={p}{{oll}}")
    line("oltl", "Other long-term liabilities", "=" + h("other_lt_liabilities"), lambda y, c, p: f"={p}{{oltl}}")
    line("tl", "Total liabilities", "=SUM(C{accounts_payable}:C{oltl})", lambda y, c, p: f"=SUM({c}{{accounts_payable}}:{c}{{oltl}})", bold=True)
    line("equity", "Total equity", "=" + h("total_equity"),
         lambda y, c, p: f"={p}{{equity}}+{c}{{ni}}+{c}{{sbc}}-{c}{{div}}-{c}{{buyback}}")
    line("tle", "Total liabilities and equity", "=C{tl}+C{equity}", lambda y, c, p: f"={c}{{tl}}+{c}{{equity}}", bold=True)
    line("check", "Balance check (assets - liabilities - equity)", "=ROUND(C{ta}-C{tle},3)", lambda y, c, p: f"=ROUND({c}{{ta}}-{c}{{tle}},3)", NUM1)
    line("nwc", "Operating working capital", "=SUM(C{receivables}:C{other_current_assets})-SUM(C{accounts_payable}:C{dmf})",
         lambda y, c, p: f"=SUM({c}{{receivables}}:{c}{{other_current_assets}})-SUM({c}{{accounts_payable}}:{c}{{dmf}})")
    section("Cash flow statement")
    line("cf_ni", "Net income", None, lambda y, c, p: f"={c}{{ni}}")
    line("cf_da", "D&A", None, lambda y, c, p: f"={c}{{da}}")
    line("sbc", "Stock-based compensation", "=" + h("sbc"), lambda y, c, p: f"={a('sbc_pct', y)}*{c}{{revenue}}")
    line("dnwc", "Decrease / (increase) in working capital", None, lambda y, c, p: f"=-({c}{{nwc}}-{p}{{nwc}})")
    line("cfo", "Cash from operations", "=" + h("cfo"), lambda y, c, p: f"={c}{{cf_ni}}+{c}{{cf_da}}+{c}{{sbc}}+{c}{{dnwc}}", bold=True)
    line("capex", "Capital expenditure", "=" + h("capex"),
         lambda y, c, p: (f"={ref('Assumptions', 'C' + str(Ar['capex_fy25']))}" if y == 2025 else f"={a('capex_pct', y)}*{c}{{revenue}}"))
    line("cfi", "Cash from investing", None, lambda y, c, p: f"=-{c}{{capex}}")
    line("div", "Dividends", "=" + h("dividends_paid"), lambda y, c, p: f"={a('payout', y)}*{c}{{ni}}")
    line("buyback", "Share repurchases", "=" + h("buybacks"), lambda y, c, p: f"={a('buybacks', y)}")
    line("cff", "Cash from financing", None, lambda y, c, p: f"=-{c}{{div}}-{c}{{buyback}}")
    line("net_change", "Net change in cash", None, lambda y, c, p: f"={c}{{cfo}}+{c}{{cfi}}+{c}{{cff}}", bold=True)
    line("fcf", "Free cash flow (CFO - capex)", "=C{cfo}-C{capex}", lambda y, c, p: f"={c}{{cfo}}-{c}{{capex}}")
    for row in s.ws.iter_rows(min_row=6):
        for c in row:
            c.value = fill_placeholders(c.value, R)
    for y in [2024] + FC_YEARS:
        s.ws[f"{cols[y]}{R['check']}"].fill = FILL_KEY
    s.widths(A=44, B=6, **{cols[y]: 12 for y in [2024] + FC_YEARS})
    s.ws.freeze_panes = "C6"
    return s, cols


def build_beta(wb):
    s = S(wb.create_sheet("Beta"))
    s.title("Beta regression: Costco vs. S&P 500, 60 monthly total returns (Feb 2020 - Jan 2025)",
            "Source: Yahoo Finance dividend/split-adjusted monthly closes (data/raw/market/yahoo_*_monthly.csv).")
    s.header(5, ["Month", "COST adj. close", "S&P 500 close", "COST return", "S&P 500 return"])
    c = prices("COST")["adjclose"]
    m = prices("^GSPC")["adjclose"]
    c = c[c.index <= "2025-01-31"].iloc[-61:]
    m = m.loc[c.index]
    r0 = 6
    for i, (d, cv) in enumerate(c.items()):
        r = r0 + i
        s.cell(f"A{r}", d.strftime("%b %Y"))
        s.cell(f"B{r}", float(cv), font=F_IN, fmt="#,##0.00")
        s.cell(f"C{r}", float(m[d]), font=F_IN, fmt="#,##0.00")
        if i:
            s.cell(f"D{r}", f"=B{r}/B{r - 1}-1", fmt=PCT2)
            s.cell(f"E{r}", f"=C{r}/C{r - 1}-1", fmt=PCT2)
    last = r0 + len(c) - 1
    s.cell("G5", "Regression", font=F_BOLD)
    s.cell("G6", "Raw beta (SLOPE)")
    s.cell("H6", f"=SLOPE(D{r0 + 1}:D{last},E{r0 + 1}:E{last})", fmt=DEC)
    s.cell("G7", "R-squared (RSQ)")
    s.cell("H7", f"=RSQ(D{r0 + 1}:D{last},E{r0 + 1}:E{last})", fmt=DEC)
    s.cell("G8", "Observations")
    s.cell("H8", f"=COUNT(D{r0 + 1}:D{last})", fmt=NUM)
    s.widths(A=11, B=15, C=15, D=13, E=15, F=3, G=20, H=10)
    return s


def build_wacc(wb, A):
    s = S(wb.create_sheet("WACC"))
    s.title("Weighted average cost of capital, 31 Jan 2025", "WACC = E/(D+E) x Re + D/(D+E) x Rd x (1 - T);  Re = Rf + beta x ERP")
    q = pd.read_csv(PROCESSED / "quarter_q1_fy2025.csv")
    qv = lambda k: float(q[(q.line_item == k) & (q.period == "Q1 FY2025")].value.iloc[0])
    fl = fin_long()
    fin_lease = float(fl[(fl.line_item == "finance_lease_liab_total") & (fl.fiscal_year == 2024)].value.iloc[0])
    om = pd.read_csv(PROCESSED / "operating_metrics_long.csv").set_index("metric")["value"]
    from src.peers import shares_outstanding
    sh_basic, sh_date = shares_outstanding("COST")
    spec = [
        ("rf", "Risk-free rate (10-year Treasury, 31 Jan 2025)", 0.0458, PCT2, "FRED DGS10 = 4.58% on 31 Jan 2025 (data/raw/market/fred_DGS10.csv)"),
        ("erp", "Equity risk premium", 0.0433, PCT2, "Damodaran implied ERP (FCFE), start of 2025 = 4.33% (histimpl.xls)"),
        ("raw_beta", "Raw beta", "=Beta!$H$6", DEC, "Beta sheet: SLOPE of 60 monthly returns"),
        ("blume_w", "Blume weight on raw beta", 0.67, DEC, "Adjusted beta = 0.67 x raw + 0.33 x 1.0"),
        ("beta", "Adjusted beta", "=B{blume_w}*B{raw_beta}+(1-B{blume_w})", DEC, None),
        ("ke", "Cost of equity", "=B{rf}+B{beta}*B{erp}", PCT2, None),
        ("aa", "ICE BofA AA corporate OAS", 0.0046, PCT2, "FRED BAMLC0A2CAA, 31 Jan 2025"),
        ("a", "ICE BofA A corporate OAS", 0.0069, PCT2, "FRED BAMLC0A3CA, 31 Jan 2025"),
        ("kd", "Pre-tax cost of debt", "=B{rf}+AVERAGE(B{aa},B{a})", PCT2, None),
        ("tax", "Tax rate", "=" + ref("Assumptions", f"C{A.rows['tax']}"), PCT, None),
        ("kd_at", "After-tax cost of debt", "=B{kd}*(1-B{tax})", PCT2, None),
        ("price", "Share price, 31 Jan 2025 ($)", 979.88, USD2, "Yahoo Finance close, 31 Jan 2025"),
        ("sh_basic", "Basic shares outstanding (m)", round(sh_basic, 5), NUM1, f"10-Q cover page, shares outstanding at {sh_date}"),
        ("rsu", "RSU dilution (m)", round(qv("shares_diluted") - qv("shares_basic"), 3), DEC,
         "Q1 FY2025 diluted less basic weighted shares (treasury-stock method)"),
        ("sh_dil", "Diluted shares (m)", "=B{sh_basic}+B{rsu}", NUM1, None),
        ("mcap", "Equity market value ($m)", "=B{price}*B{sh_dil}", NUM, None),
        ("debt_fv", "Fair value of long-term debt ($m)", float(om["q1_debt_fair_value"]), NUM,
         "Q1 FY2025 10-Q: fair value of long-term debt incl. current portion"),
        ("fin_lease", "Finance lease liabilities ($m)", fin_lease, NUM, "FY2024 10-K lease note: total lease liabilities 4,052 less operating 2,554"),
        ("d_mv", "Debt incl. finance leases ($m)", "=B{debt_fv}+B{fin_lease}", NUM, None),
        ("we", "Weight of equity", "=B{mcap}/(B{mcap}+B{d_mv})", PCT2, None),
        ("wd", "Weight of debt", "=B{d_mv}/(B{mcap}+B{d_mv})", PCT2, None),
        ("wacc", "WACC", "=B{we}*B{ke}+B{wd}*B{kd_at}", PCT2, None),
    ]
    s.header(5, ["Item", "Value", "Source / note"])
    for i, (key, *_rest) in enumerate(spec):
        s.rows[key] = 6 + i
    for key, lab, val, fmt, note in spec:
        rr = s.rows[key]
        v = fill_placeholders(val, s.rows)
        formula = isinstance(v, str) and v.startswith("=")
        font = F_IN if not formula else (F_LINK if "!" in v and not any(o in v[1:] for o in "+-*/(") else F_CALC)
        s.cell(f"A{rr}", lab, bold=key in ("ke", "wacc"))
        s.cell(f"B{rr}", v, font=font, fmt=fmt, fill=FILL_KEY if key in ("wacc", "erp", "rf") else None, bold=key in ("ke", "wacc"))
        if note:
            s.cell(f"C{rr}", note, font=F_NOTE)
    s.widths(A=46, B=14, C=90)
    return s


def build_dcf(wb, F, fcols, W, A):
    s = S(wb.create_sheet("DCF"))
    s.title("Discounted cash flow valuation, 31 Jan 2025 ($ millions)",
            "UFCF = EBIT x (1 - T) + D&A - capex - increase in working capital. Mid-year convention; FY2025 is a stub from the 24 Nov 2024 balance sheet.")
    Fr, Wr, Ar = F.rows, W.rows, A.rows
    w = lambda k: ref("WACC", f"B{Wr[k]}")
    cols = {y: L(3 + i) for i, y in enumerate(FC_YEARS)}   # C..G
    s.cell("A5", "Valuation date", bold=True)
    s.cell("B5", date(2025, 1, 31), font=F_IN, fmt=DATE)
    s.cell("A6", "Balance-sheet date (Q1 FY2025 10-Q)", bold=True)
    s.cell("B6", date(2024, 11, 24), font=F_IN, fmt=DATE)
    s.cell("A7", "FY2025 start", bold=True)
    s.cell("B7", date(2024, 9, 2), font=F_IN, fmt=DATE)
    s.cell("A8", "WACC", bold=True)
    s.cell("B8", "=" + w("wacc"), font=F_LINK, fmt=PCT2, fill=FILL_KEY)
    s.cell("A9", "Terminal growth (g)", bold=True)
    s.cell("B9", "=" + ref("Assumptions", f"C{Ar['tg']}"), font=F_LINK, fmt=PCT, fill=FILL_KEY)
    s.cell("A10", "Terminal RONIC", bold=True)
    s.cell("B10", "=" + ref("Assumptions", f"C{Ar['ronic']}"), font=F_LINK, fmt=PCT, fill=FILL_KEY)
    s.cell("A11", "Tax rate", bold=True)
    s.cell("B11", "=" + w("tax"), font=F_LINK, fmt=PCT)
    s.header(13, ["$ millions", ""] + [f"FY{y}E" for y in FC_YEARS])
    fy_end = {2025: date(2025, 8, 31), 2026: date(2026, 8, 30), 2027: date(2027, 8, 29), 2028: date(2028, 9, 3), 2029: date(2029, 9, 2)}
    R = s.rows
    f = lambda key, y: "=" + ref("Forecast", f"{fcols[y]}{Fr[key]}", absolute=False)
    prev_col = lambda y: L(2 + FC_YEARS.index(y))
    spec = [
        ("fy_end", "Fiscal year end", lambda y, c: fy_end[y], DATE),
        ("ebit", "EBIT", lambda y, c: f("ebit", y), NUM),
        ("tax_ebit", "Taxes on EBIT", lambda y, c: f"={c}{{ebit}}*$B$11", NUM),
        ("nopat", "NOPAT", lambda y, c: f"={c}{{ebit}}-{c}{{tax_ebit}}", NUM),
        ("da", "+ D&A", lambda y, c: f("da", y), NUM),
        ("capex", "- Capex", lambda y, c: f("capex", y), NUM),
        ("dnwc", "+ Working capital released", lambda y, c: f("dnwc", y), NUM),
        ("ufcf", "Unlevered free cash flow", lambda y, c: f"={c}{{nopat}}+{c}{{da}}-{c}{{capex}}+{c}{{dnwc}}", NUM),
        ("frac", "Share of year counted", lambda y, c: ("=(C{fy_end}-$B$6)/(C{fy_end}-$B$7+1)" if y == 2025 else 1), "0.000"),
        ("start", "Period start", lambda y, c: ("=$B$6" if y == 2025 else f"={prev_col(y)}{{fy_end}}"), DATE),
        ("mid", "Period midpoint", lambda y, c: f"={c}{{start}}+INT(({c}{{fy_end}}-{c}{{start}})/2)", DATE),
        ("t", "Discount time (years from valuation date)", lambda y, c: f"=MAX(({c}{{mid}}-$B$5)/365.25,0)", "0.000"),
        ("cf", "UFCF in period", lambda y, c: f"={c}{{ufcf}}*{c}{{frac}}", NUM),
        ("df", "Discount factor", lambda y, c: f"=1/(1+$B$8)^{c}{{t}}", "0.0000"),
        ("pv", "PV of UFCF", lambda y, c: f"={c}{{cf}}*{c}{{df}}", NUM),
        ("ebitda", "EBITDA", lambda y, c: f("ebitda", y), NUM),
    ]
    for i, (key, *_r) in enumerate(spec):
        R[key] = 14 + i
    for key, lab, fn, fmt in spec:
        rr = R[key]
        s.cell(f"A{rr}", lab, bold=key in ("ufcf", "pv"))
        for y in FC_YEARS:
            v = fill_placeholders(fn(y, cols[y]), R)
            font = F_IN if key in ("fy_end",) or (key == "frac" and y != 2025) else (F_LINK if isinstance(v, str) and "Forecast!" in v else F_CALC)
            s.cell(f"{cols[y]}{rr}", v, font=font, fmt=fmt, bold=key in ("ufcf", "pv"))
    G = cols[2029]
    r = 14 + len(spec) + 1
    s.cell(f"A{r}", "Terminal value", font=F_BOLD, fill=FILL_SUB)
    s.cell(f"B{r}", None, fill=FILL_SUB)
    r += 1
    tail = [
        ("nopat_next", "NOPAT FY2030 = EBIT FY2029 x (1 + g) x (1 - T)", f"={G}{{ebit}}*(1+$B$9)*(1-$B$11)", NUM, None),
        ("reinv", "Reinvestment rate = g / RONIC", "=$B$9/$B$10", PCT, None),
        ("fcf_next", "Terminal FCF FY2030", "=B{nopat_next}*(1-B{reinv})", NUM, None),
        ("tv", "Terminal value = FCF / (WACC - g)", "=B{fcf_next}/($B$8-$B$9)", NUM, None),
        ("pv_tv", "PV of terminal value (discounted consistently with mid-year flows)", f"=B{{tv}}/(1+$B$8)^{G}{{t}}", NUM, None),
        (None, "Enterprise value to equity value", None, None, None),
        ("sum_pv", "Sum of PV of UFCF", f"=SUM({cols[2025]}{{pv}}:{G}{{pv}})", NUM, None),
        ("pv_tv2", "PV of terminal value", "=B{pv_tv}", NUM, None),
        ("ev", "Enterprise value", "=B{sum_pv}+B{pv_tv2}", NUM, None),
        ("tv_share", "Terminal value % of EV", "=B{pv_tv2}/B{ev}", PCT, None),
    ]
    q = pd.read_csv(PROCESSED / "quarter_q1_fy2025.csv")
    qv = lambda k: float(q[(q.line_item == k) & (q.period == "Q1 FY2025")].value.iloc[0])
    tail += [
        ("cash", "+ Cash and cash equivalents (24 Nov 2024)", qv("cash"), NUM, "Q1 FY2025 10-Q balance sheet"),
        ("sti", "+ Short-term investments", qv("short_term_investments"), NUM, "Q1 FY2025 10-Q balance sheet"),
        ("debt", "- Debt (current + long-term, carrying value)", qv("current_debt") + qv("long_term_debt"), NUM, "Q1 FY2025 10-Q: 97 + 5,745"),
        ("fl", "- Finance lease liabilities", "=" + w("fin_lease"), NUM, None),
        ("equity", "Equity value", "=B{ev}+B{cash}+B{sti}-B{debt}-B{fl}", NUM, None),
        ("shares", "Diluted shares (m)", "=" + w("sh_dil"), NUM1, None),
        ("vps", "Value per share ($)", "=B{equity}/B{shares}", USD2, None),
        ("price", "Share price 31 Jan 2025 ($)", "=" + w("price"), USD2, None),
        ("updown", "Upside / (downside)", "=B{vps}/B{price}-1", PCT, None),
        ("exit_mult", "Implied terminal EV / FY2029 EBITDA (mid-year basis)", f"=B{{tv}}/{G}{{ebitda}}", MULT, None),
        ("exit_mult_ye", "Implied terminal EV / FY2029 EBITDA (FY2029 year-end basis; compare with trading multiples)",
         f"=B{{tv}}*(1+$B$8)^0.5/{G}{{ebitda}}", MULT, None),
        ("net_cash", "Net cash", "=B{cash}+B{sti}-B{debt}-B{fl}", NUM, None),
        ("mkt_ev", "Market EV (market cap - net cash)", "=" + w("mcap") + "-B{net_cash}", NUM, None),
        ("mkt_mult29", "Market EV / FY2029E EBITDA", f"=B{{mkt_ev}}/{G}{{ebitda}}", MULT, None),
        ("mkt_mult", "Market EV / FY2025E EBITDA", f"=B{{mkt_ev}}/{cols[2025]}{{ebitda}}", MULT, None),
        ("impl_g", "Reverse DCF: terminal growth implied by the price", 0.070525, PCT2,
         "Solved in src/dcf.py by bisection. To re-solve here: Data > What-If Analysis > Goal Seek, set DCF value per share = price by changing Assumptions!E(terminal g, Base)."),
    ]
    for key, *_x in tail:
        if key:
            R[key] = r
        r += 1
    r = R["nopat_next"]
    for key, lab, fml, fmt, note in tail:
        if key is None:
            s.cell(f"A{r}", lab, font=F_BOLD, fill=FILL_SUB)
            s.cell(f"B{r}", None, fill=FILL_SUB)
            r += 1
            continue
        v = fill_placeholders(fml, R)
        formula = isinstance(v, str) and v.startswith("=")
        font = F_IN if not formula else (F_LINK if "WACC!" in v and not any(o in v[1:] for o in "+-*/(") else F_CALC)
        big = key in ("ev", "equity", "vps")
        s.cell(f"A{r}", lab, bold=big)
        s.cell(f"B{r}", v, font=font, fmt=fmt, bold=big, fill=FILL_KEY if key == "vps" else None)
        if note:
            s.cell(f"C{r}", note, font=F_NOTE)
        r += 1
    s.widths(A=60, B=14, C=13, D=13, E=13, F=13, G=13)
    return s, cols


def build_sensitivity(wb, D, dcols):
    s = S(wb.create_sheet("Sensitivity"))
    s.title("Sensitivity analysis: DCF value per share ($)",
            "Grids 1-2 are live formulas that recompute the DCF for each cell. Grid 3 needs the forecast re-run per cell, so it is pasted from src/sensitivity.py.")
    R = D.rows
    first, last = dcols[2025], dcols[2029]
    cf = f"DCF!${first}${R['cf']}:${last}${R['cf']}"
    tt = f"DCF!${first}${R['t']}:${last}${R['t']}"
    ebit29 = f"DCF!${last}${R['ebit']}"
    t_last = f"DCF!${last}${R['t']}"
    tax, ronic, w0 = "DCF!$B$11", "DCF!$B$10", "DCF!$B$8"
    netcash, shares = f"DCF!$B${R['net_cash']}", f"DCF!$B${R['shares']}"

    def val(wc, gc, rc):
        return (f"=(SUMPRODUCT({cf},(1+{wc})^(-{tt}))+{ebit29}*(1+{gc})*(1-{tax})*(1-{gc}/{rc})/({wc}-{gc})/(1+{wc})^{t_last}"
                f"+{netcash})/{shares}")

    s.cell("A5", "Row and column headers in blue are editable; WACC rows are centred on the live WACC. Yellow = base case.", font=F_NOTE)
    gs = [0.02, 0.025, 0.03, 0.035, 0.04]
    r = 7
    s.cell(f"A{r}", "1. WACC x terminal growth", font=F_BOLD)
    s.cell(f"A{r + 1}", "WACC (rows) x terminal growth (columns); RONIC held at base", font=F_NOTE)
    for j, g in enumerate(gs):
        s.cell(f"{L(2 + j)}{r + 2}", g, font=F_IN, fmt=PCT, fill=FILL_SUB)
    for i, off in enumerate([-0.01, -0.005, 0.0, 0.005, 0.01]):
        rr = r + 3 + i
        s.cell(f"A{rr}", f"={w0}{off:+.3f}" if off else f"={w0}", fmt=PCT2, fill=FILL_SUB)
        for j in range(5):
            s.cell(f"{L(2 + j)}{rr}", val(f"$A{rr}", f"{L(2 + j)}${r + 2}", ronic), fmt=USD, fill=FILL_KEY if (i == 2 and j == 2) else None)
    s.rows["wacc_g_centre"] = r + 5
    r += 9
    s.cell(f"A{r}", "2. Terminal RONIC x terminal growth (terminal-year cash flow)", font=F_BOLD)
    s.cell(f"A{r + 1}", "RONIC (rows) x terminal growth (columns); WACC held at base", font=F_NOTE)
    for j, g in enumerate(gs):
        s.cell(f"{L(2 + j)}{r + 2}", g, font=F_IN, fmt=PCT, fill=FILL_SUB)
    for i, rn in enumerate([0.15, 0.20, 0.25, 0.30, 0.35]):
        rr = r + 3 + i
        s.cell(f"A{rr}", rn, font=F_IN, fmt=PCT, fill=FILL_SUB)
        for j in range(5):
            s.cell(f"{L(2 + j)}{rr}", val(w0, f"{L(2 + j)}${r + 2}", f"$A{rr}"), fmt=USD, fill=FILL_KEY if (i == 2 and j == 2) else None)
    s.rows["ronic_g_centre"] = r + 5
    r += 9
    s.cell(f"A{r}", "3. Comparable-sales growth x gross margin (static values from src/sensitivity.py)", font=F_BOLD)
    s.cell(f"A{r + 1}", "Each cell re-runs the full three-statement forecast, which a formula grid cannot do; re-run python -m src.sensitivity to refresh.",
           font=F_NOTE)
    rm = pd.read_csv(TABLES / "sens_revenue_margin.csv", index_col=0)
    for j, cname in enumerate(rm.columns):
        s.cell(f"{L(2 + j)}{r + 2}", cname.replace("gross margin ", "GM "), fill=FILL_SUB)
    for i, (idx, row) in enumerate(rm.iterrows()):
        rr = r + 3 + i
        s.cell(f"A{rr}", idx, fill=FILL_SUB)
        for j, v in enumerate(row.values):
            s.cell(f"{L(2 + j)}{rr}", float(v), font=F_IN, fmt=USD, fill=FILL_KEY if (i == 2 and j == 2) else None)
    r += 9
    s.cell(f"A{r}", "4. Bear / base / bull (static values from src/sensitivity.py; set Assumptions!C5 to 1 or 3 to run a case live)", font=F_BOLD)
    scn = pd.read_csv(TABLES / "sens_scenarios.csv", index_col=0)
    s.header(r + 1, ["", "Bear", "Base", "Bull"])
    items = [("revenue_cagr_fy24_29", "FY24-29 revenue CAGR", PCT), ("operating_margin_fy2029", "FY2029 operating margin", PCT2),
             ("terminal_growth", "Terminal growth", PCT), ("terminal_ronic", "Terminal RONIC", PCT),
             ("value_per_share", "Value per share ($)", USD), ("probability", "Probability", PCT)]
    for i, (k, lab, fmt) in enumerate(items):
        rr = r + 2 + i
        s.cell(f"A{rr}", lab)
        for j, name in enumerate(["bear", "base", "bull"]):
            s.cell(f"{L(2 + j)}{rr}", float(scn.loc[name, k]), font=F_IN, fmt=fmt)
        s.rows[f"scn_{k}"] = rr
    pw = r + 2 + len(items)
    s.cell(f"A{pw}", "Probability-weighted value ($)", bold=True)
    s.cell(f"B{pw}", f"=SUMPRODUCT(B{s.rows['scn_value_per_share']}:D{s.rows['scn_value_per_share']},"
                     f"B{s.rows['scn_probability']}:D{s.rows['scn_probability']})", fmt=USD, bold=True)
    s.rows["prob_weighted"] = pw
    s.widths(A=44, B=13, C=13, D=13, E=13, F=13)
    return s


def build_comps(wb):
    s = S(wb.create_sheet("Comps"))
    s.title("Comparable-company analysis, 31 Jan 2025 ($ millions)",
            "Inputs (blue) from SEC XBRL company facts filed on or before 31 Jan 2025 and Yahoo Finance; LTM on a 52-week basis (src/ltm.py). "
            "EV excludes operating leases.")
    df = pd.read_csv(TABLES / "comps.csv", index_col=0)
    peers = json.load(open(CONFIG_DIR / "peers.json"))
    tickers = list(df.index)
    cols = {t: L(3 + i) for i, t in enumerate(tickers)}      # C onwards, one column per company
    s.header(5, ["", "Unit"] + tickers)
    R = s.rows
    inputs = [("price", "Share price 31 Jan 2025", "$", USD2), ("shares_m", "Shares outstanding (cover page)", "m", NUM1),
              ("financial_debt", "Debt incl. finance leases", "$m", NUM), ("noncontrolling_interest", "Noncontrolling interest", "$m", NUM),
              ("cash_and_sti", "Cash and short-term investments", "$m", NUM), ("ltm_revenue", "LTM revenue", "$m", NUM),
              ("ltm_ebit", "LTM operating income", "$m", NUM), ("ltm_d_and_a", "LTM D&A", "$m", NUM),
              ("ltm_net_income", "LTM net income", "$m", NUM), ("revenue_growth_ytd", "Revenue growth, YTD y/y", "%", PCT),
              ("revenue_cagr_3y", "Revenue CAGR, 3 years", "%", PCT)]
    r = 6
    for key, lab, unit, fmt in inputs:
        s.cell(f"A{r}", lab)
        s.cell(f"B{r}", unit)
        for t in tickers:
            note = None
            if key == "financial_debt":
                note = f"Balance sheet {df.loc[t, 'bs_date']}; finance leases: {df.loc[t, 'finance_lease_source']}"
            elif key == "ltm_revenue":
                note = f"LTM to {df.loc[t, 'ltm_end']} (fiscal year + YTD - prior YTD), scaled to 52 weeks"
            s.cell(f"{cols[t]}{r}", float(df.loc[t, key]), font=F_IN, fmt=fmt, comment=note)
        R[key] = r
        r += 1
    r += 1
    calc = [("mcap", "Market capitalisation", "={c}{price}*{c}{shares_m}", NUM),
            ("ev", "Enterprise value", "={c}{mcap}+{c}{financial_debt}+{c}{noncontrolling_interest}-{c}{cash_and_sti}", NUM),
            ("ebitda", "LTM EBITDA", "={c}{ltm_ebit}+{c}{ltm_d_and_a}", NUM),
            ("ebitda_m", "EBITDA margin", "={c}{ebitda}/{c}{ltm_revenue}", PCT),
            ("ev_rev", "EV / Revenue", "={c}{ev}/{c}{ltm_revenue}", MULT),
            ("ev_ebitda", "EV / EBITDA", '=IF({c}{ebitda}>0,{c}{ev}/{c}{ebitda},"NM")', MULT),
            ("pe", "P / E", '=IF({c}{ltm_net_income}>0,{c}{mcap}/{c}{ltm_net_income},"NM")', MULT),
            ("ps", "P / S", "={c}{mcap}/{c}{ltm_revenue}", MULT),
            ("gadj", "Growth-adjusted EV/EBITDA", '=IF(ISNUMBER({c}{ev_ebitda}),{c}{ev_ebitda}/({c}{revenue_cagr_3y}*100),"NM")', "0.00")]
    for i, (key, *_x) in enumerate(calc):
        R[key] = r + i
    for key, lab, fml, fmt in calc:
        s.cell(f"A{R[key]}", lab, bold=key in ("ev_ebitda", "pe"))
        for t in tickers:
            s.cell(f"{cols[t]}{R[key]}", fill_placeholders(fml.replace("{c}", cols[t]), R), fmt=fmt, bold=key in ("ev_ebitda", "pe"),
                   align="right")
    r += len(calc)
    s.cell(f"A{r}", "Tier", font=F_NOTE)
    for t in tickers:
        s.cell(f"{cols[t]}{r}", "Subject" if t == "COST" else "Tier " + peers["tiers"][t].split(" - ")[0], font=F_NOTE, align="center")
    r += 2
    rng = lambda key: f"{cols[tickers[1]]}{R[key]}:{cols[tickers[-1]]}{R[key]}"
    tier1 = lambda key: ",".join(f"{cols[t]}{R[key]}" for t in tickers[1:] if peers["tiers"][t].startswith("1"))
    s.header(r, ["Peer statistics (excl. Costco)", "", "25th pct", "Median", "75th pct", "Tier 1 median", "Costco", "vs. median",
                 "Walmart"])
    r += 1
    for key, lab in [("ev_rev", "EV / Revenue"), ("ev_ebitda", "EV / EBITDA"), ("pe", "P / E"), ("ps", "P / S"),
                     ("gadj", "Growth-adjusted EV/EBITDA")]:
        s.cell(f"A{r}", lab)
        s.cell(f"C{r}", f"=QUARTILE({rng(key)},1)", fmt=MULT)
        s.cell(f"D{r}", f"=MEDIAN({rng(key)})", fmt=MULT, bold=True)
        s.cell(f"E{r}", f"=QUARTILE({rng(key)},3)", fmt=MULT)
        s.cell(f"F{r}", f"=MEDIAN({tier1(key)})", fmt=MULT)
        s.cell(f"G{r}", f"={cols['COST']}{R[key]}", fmt=MULT)
        s.cell(f"H{r}", f"=G{r}/D{r}-1", fmt=PCT)
        s.cell(f"I{r}", f"={cols['WMT']}{R[key]}", fmt=MULT)
        R[f"stat_{key}"] = r
        r += 1
    r += 1
    s.header(r, ["Implied Costco value per share ($)", "Costco metric", "25th pct", "Median", "75th pct", "Tier 1 median", "Walmart"])
    r += 1
    c = cols["COST"]
    netcash = f"({c}{R['cash_and_sti']}-{c}{R['financial_debt']}-{c}{R['noncontrolling_interest']})"
    for key, lab, metric, is_ev in [("ev_ebitda", "EV / EBITDA", "ebitda", True), ("ev_rev", "EV / Revenue", "ltm_revenue", True),
                                    ("pe", "P / E", "ltm_net_income", False), ("ps", "P / S", "ltm_revenue", False)]:
        s.cell(f"A{r}", lab)
        s.cell(f"B{r}", f"={c}{R[metric]}", fmt=NUM)
        for col, scol in zip("CDEFG", "CDEFI"):       # G = Walmart alone (statistic in column I)
            m = f"{scol}{R['stat_' + key]}"
            s.cell(f"{col}{r}", f"=({m}*$B{r}{'+' + netcash if is_ev else ''})/{c}{R['shares_m']}", fmt=USD, bold=col == "D")
        R[f"impl_{key}"] = r
        r += 1
    s.widths(A=38, B=14, **{**{cols[t]: 13 for t in tickers}, L(3 + len(tickers)): 16})
    return s


def build_summary(wb, D, SE, C):
    s = S(wb.create_sheet("Summary"))
    s.title("Valuation summary ($ per share)", "Football-field inputs; every number links to the sheet that computes it.")
    s.header(5, ["Method", "Low", "Central", "High"])
    Dr, Cr, SEr = D.rows, C.rows, SE.rows
    wc = SEr["wacc_g_centre"]
    vps = SEr["scn_value_per_share"]
    rows = [
        ("DCF: bear / base / bull", f"=Sensitivity!B{vps}", f"=DCF!B{Dr['vps']}", f"=Sensitivity!D{vps}"),
        ("DCF: WACC +/-0.5pp, growth +/-0.5pp", f"=MIN(Sensitivity!C{wc - 1}:E{wc + 1})", f"=DCF!B{Dr['vps']}", f"=MAX(Sensitivity!C{wc - 1}:E{wc + 1})"),
        ("Comps: EV/EBITDA (peer 25th-75th)", f"=Comps!C{Cr['impl_ev_ebitda']}", f"=Comps!D{Cr['impl_ev_ebitda']}", f"=Comps!E{Cr['impl_ev_ebitda']}"),
        ("Comps: P/E (peer 25th-75th)", f"=Comps!C{Cr['impl_pe']}", f"=Comps!D{Cr['impl_pe']}", f"=Comps!E{Cr['impl_pe']}"),
        ("Comps: EV/Revenue (peer 25th-75th)", f"=Comps!C{Cr['impl_ev_rev']}", f"=Comps!D{Cr['impl_ev_rev']}", f"=Comps!E{Cr['impl_ev_rev']}"),
        ("Comps selected range: Tier 1 median EV/EBITDA to Walmart P/E", f"=Comps!F{Cr['impl_ev_ebitda']}", None, f"=Comps!G{Cr['impl_pe']}"),
    ]
    r = 6
    for lab, lo, mid, hi in rows:
        s.cell(f"A{r}", lab)
        for col, v in zip("BCD", [lo, mid, hi]):
            if v:
                s.cell(f"{col}{r}", v, font=F_LINK if v.count("!") == 1 and not v.startswith("=M") else F_CALC, fmt=USD)
        r += 1
    r += 1
    s.cell(f"A{r}", "Share price, 31 Jan 2025", bold=True)
    s.cell(f"C{r}", f"=DCF!B{Dr['price']}", font=F_LINK, fmt=USD2, bold=True)
    s.cell(f"A{r + 1}", "Probability-weighted DCF value", bold=True)
    s.cell(f"C{r + 1}", f"=Sensitivity!B{SEr['prob_weighted']}", font=F_LINK, fmt=USD, bold=True)
    s.cell(f"A{r + 2}", "DCF value vs. price", bold=True)
    s.cell(f"C{r + 2}", f"=DCF!B{Dr['updown']}", font=F_LINK, fmt=PCT, bold=True)
    s.widths(A=46, B=12, C=12, D=12)
    return s


def build_checks(wb, F, W, D, SE, C):
    s = S(wb.create_sheet("Checks"))
    s.title("Model checks", "Integrity checks and a reconciliation of every headline number to the Python pipeline (outputs/tables). All flags should read OK.")
    s.header(5, ["Check", "Excel", "Python / expected", "Difference", "Status"])
    with contextlib.redirect_stdout(io.StringIO()):
        from src.dcf import main as dcf_main
        from src.wacc import compute
        r_ = dcf_main()
        w_ = compute()
    fc = pd.read_csv(TABLES / "forecast_income_statement.csv", index_col=0)
    comps = pd.read_csv(TABLES / "comps.csv", index_col=0)
    cstats = pd.read_csv(TABLES / "comps_statistics.csv", index_col=0)
    cimpl = pd.read_csv(TABLES / "comps_implied_value.csv", index_col=[0, 1])
    Fr, Dr, Wr, Cr = F.rows, D.rows, W.rows, C.rows
    items = [
        ("Balance sheet balances, every year (sum of |check|)", f"=SUMPRODUCT(ABS(Forecast!C{Fr['check']}:H{Fr['check']}))", 0.0, 0.01),
        ("Cash flow ties to balance-sheet cash, FY2025-29",
         f"=SUMPRODUCT(ABS(Forecast!D{Fr['cash']}:H{Fr['cash']}-Forecast!C{Fr['cash']}:G{Fr['cash']}-Forecast!D{Fr['net_change']}:H{Fr['net_change']}))",
         0.0, 0.01),
        ("FY2024 revenue links to Historicals ($m)", f"=Forecast!C{Fr['revenue']}", float(fc.loc["total_revenue", "FY2024A"]), 0.5),
        ("FY2029 revenue ($m)", f"=Forecast!H{Fr['revenue']}", float(fc.loc["total_revenue", "FY2029E"]), 0.5),
        ("FY2029 operating income ($m)", f"=Forecast!H{Fr['ebit']}", float(fc.loc["operating_income", "FY2029E"]), 0.5),
        ("FY2025 EPS ($)", f"=Forecast!D{Fr['eps']}", float(fc.loc["eps_diluted", "FY2025E"]), 0.005),
        ("Raw beta", "=Beta!H6", float(w_["raw_beta"]), 0.0005),
        ("WACC", f"=WACC!B{Wr['wacc']}", float(w_["wacc"]), 0.00005),
        ("Enterprise value ($m)", f"=DCF!B{Dr['ev']}", float(r_["enterprise_value"]), 1.0),
        ("Value per share ($)", f"=DCF!B{Dr['vps']}", float(r_["value_per_share"]), 0.05),
        ("Terminal value % of EV", f"=DCF!B{Dr['tv_share']}", float(r_["tv_share_of_ev"]), 0.0005),
        ("Implied exit multiple, FY2029 year-end basis", f"=DCF!B{Dr['exit_mult_ye']}", float(r_["implied_tv_ev_ebitda_fy29_yearend"]), 0.005),
        ("Sensitivity grid centre = base value ($)", f"=Sensitivity!D{SE.rows['wacc_g_centre']}", float(r_["value_per_share"]), 0.05),
        ("Comps: Costco EV/EBITDA", f"=Comps!G{Cr['stat_ev_ebitda']}", float(comps.loc["COST", "ev_ebitda"]), 0.005),
        ("Comps: peer median EV/EBITDA", f"=Comps!D{Cr['stat_ev_ebitda']}", float(cstats.loc["ev_ebitda", "peer_median"]), 0.005),
        ("Comps: selected range, Walmart P/E implied value ($)", f"=Comps!G{Cr['impl_pe']}",
         float(cimpl.loc[("pe", "walmart"), "implied_price"]), 0.05),
        ("Scenario selector is on Base (2)", "=Assumptions!C5", 2, 0.0),
    ]
    r = 6
    for lab, fml, expected, tol in items:
        s.cell(f"A{r}", lab)
        s.cell(f"B{r}", fml, font=F_LINK if fml.count("!") == 1 and "(" not in fml else F_CALC, fmt="#,##0.0000")
        s.cell(f"C{r}", expected, font=F_IN, fmt="#,##0.0000")
        s.cell(f"D{r}", f"=B{r}-C{r}", fmt="#,##0.0000")
        s.cell(f"E{r}", f'=IF(ABS(D{r})<={tol},"OK","CHECK")', bold=True)
        r += 1
    s.cell(f"A{r + 1}", "All checks", bold=True)
    s.cell(f"E{r + 1}", f'=IF(COUNTIF(E6:E{r - 1},"OK")={len(items)},"ALL OK","REVIEW")', bold=True, fill=FILL_KEY)
    s.rows["all"] = r + 1
    s.widths(A=52, B=16, C=18, D=14, E=10)
    return s


def build_cover(wb, D, CH):
    s = S(wb["Cover"])
    s.cell("A1", "Costco Wholesale Corporation (COST): Equity Research and Valuation Model", font=F_TITLE)
    s.cell("A2", "Independent academic research and valuation case study; not investment advice.",
           font=Font(name=FONT, size=10, bold=True, color="C00000"))
    s.cell("A3", "Valuation date: 31 January 2025. Only information public on that date is used. Author: Pooja Master.", font=F_NOTE)
    s.cell("A5", "Headline outputs", font=F_BOLD, fill=FILL_SUB)
    s.cell("B5", None, fill=FILL_SUB)
    Dr = D.rows
    outputs = [("Active scenario", "=Assumptions!D5", None), ("WACC", "=DCF!B8", PCT2), ("Terminal growth", "=DCF!B9", PCT),
               ("Enterprise value ($m)", f"=DCF!B{Dr['ev']}", NUM), ("DCF value per share ($)", f"=DCF!B{Dr['vps']}", USD2),
               ("Share price, 31 Jan 2025 ($)", f"=DCF!B{Dr['price']}", USD2), ("Upside / (downside)", f"=DCF!B{Dr['updown']}", PCT),
               ("Model checks", f"=Checks!E{CH.rows['all']}", None)]
    for i, (lab, f, fmt) in enumerate(outputs):
        s.cell(f"A{6 + i}", lab)
        s.cell(f"B{6 + i}", f, font=F_LINK, fmt=fmt, bold=True, align="left")
    s.cell("A15", "Sheets", font=F_BOLD, fill=FILL_SUB)
    s.cell("B15", None, fill=FILL_SUB)
    guide = [("Historicals", "FY2019-FY2024 statements from 10-K XBRL; comments give the source of every number"),
             ("Assumptions", "Forecast drivers with rationale; scenario switch in C5 (1 Bear, 2 Base, 3 Bull)"),
             ("Forecast", "Integrated income statement, balance sheet and cash flow; balance check"),
             ("Beta", "60-month regression of Costco vs. S&P 500 returns"), ("WACC", "CAPM, cost of debt, market-value weights"),
             ("DCF", "UFCF, stub and mid-year timing, value-driver terminal value, equity bridge"),
             ("Sensitivity", "Live WACC x g and RONIC x g grids; revenue x margin and scenario tables"),
             ("Comps", "Peer multiples, quartiles and implied Costco value"), ("Summary", "Football-field table"),
             ("Checks", "Integrity checks and reconciliation to the Python pipeline")]
    for i, (n, d) in enumerate(guide):
        s.cell(f"A{16 + i}", n, bold=True)
        s.cell(f"B{16 + i}", d)
    s.cell("A27", "Colour legend", font=F_BOLD, fill=FILL_SUB)
    s.cell("B27", None, fill=FILL_SUB)
    s.cell("A28", "1,234", font=F_IN)
    s.cell("B28", "Blue: hardcoded input (sourced in a comment or note)")
    s.cell("A29", "1,234", font=F_CALC)
    s.cell("B29", "Black: formula")
    s.cell("A30", "1,234", font=F_LINK)
    s.cell("B30", "Green: link to another sheet")
    s.cell("A31", "1,234", fill=FILL_KEY)
    s.cell("B31", "Yellow fill: key assumption or output")
    s.cell("A33", "Built by src/build_excel.py in github.com/pooja003-cloud/costco-equity-research-valuation", font=F_NOTE)
    s.widths(A=34, B=100)


def main():
    wb = Workbook()
    wb.active.title = "Cover"
    H, hc = build_historicals(wb)
    A, ac = build_assumptions(wb)
    Fc, fcols = build_forecast(wb, H, hc, A, ac)
    build_beta(wb)
    W = build_wacc(wb, A)
    D, dcols = build_dcf(wb, Fc, fcols, W, A)
    SE = build_sensitivity(wb, D, dcols)
    C = build_comps(wb)
    build_summary(wb, D, SE, C)
    CH = build_checks(wb, Fc, W, D, SE, C)
    build_cover(wb, D, CH)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    out = OUTPUTS / "COST_Valuation_Model.xlsx"
    OUTPUTS.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print("saved", out.relative_to(OUTPUTS.parent))


if __name__ == "__main__":
    main()
